import os
import re

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
SUPPORTED_EXTENSIONS = {".csv", ".txt", ".xlsx", ".pdf", ".docx"}


class UnsupportedFileType(Exception):
    pass


def _emails_from_text(text):
    return EMAIL_RE.findall(text or "")


def _text_from_txt(f):
    return f.read().decode("utf-8", errors="ignore")


def _text_from_xlsx(f):
    import openpyxl
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    parts.append(str(cell))
    return " ".join(parts)


def _text_from_pdf(f):
    from pypdf import PdfReader
    reader = PdfReader(f)
    return " ".join((page.extract_text() or "") for page in reader.pages)


def _text_from_docx(f):
    import docx
    document = docx.Document(f)
    return " ".join(p.text for p in document.paragraphs)


_EXTRACTORS = {
    ".txt": _text_from_txt,
    ".csv": _text_from_txt,
    ".xlsx": _text_from_xlsx,
    ".pdf": _text_from_pdf,
    ".docx": _text_from_docx,
}


def extract_emails_from_file(uploaded_file):
    ext = os.path.splitext(uploaded_file.name)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(ext)
    text = _EXTRACTORS[ext](uploaded_file)
    seen = {}
    for email in _emails_from_text(text):
        seen.setdefault(email.lower(), None)
    return list(seen.keys())
