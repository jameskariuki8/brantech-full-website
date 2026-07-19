import io
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from messaging import imports
from messaging.imports import UnsupportedFileType


class ImportTests(TestCase):
    def test_txt_extraction_dedupes_and_lowercases(self):
        content = b"Reach me at Alice@Example.com or bob@x.com. Again: alice@example.com"
        f = SimpleUploadedFile("contacts.txt", content, content_type="text/plain")
        self.assertEqual(
            sorted(imports.extract_emails_from_file(f)),
            ["alice@example.com", "bob@x.com"],
        )

    def test_csv_extraction(self):
        content = b"name,email\nAda,ada@x.com\nBob,bob@x.com\n"
        f = SimpleUploadedFile("list.csv", content, content_type="text/csv")
        self.assertEqual(
            sorted(imports.extract_emails_from_file(f)), ["ada@x.com", "bob@x.com"]
        )

    def test_xlsx_extraction(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name", "email"])
        ws.append(["Ada", "ada@x.com"])
        buf = io.BytesIO()
        wb.save(buf)
        f = SimpleUploadedFile(
            "list.xlsx", buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(imports.extract_emails_from_file(f), ["ada@x.com"])

    def test_docx_extraction(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("Contact ada@x.com for details.")
        buf = io.BytesIO()
        doc.save(buf)
        f = SimpleUploadedFile(
            "list.docx", buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertEqual(imports.extract_emails_from_file(f), ["ada@x.com"])

    def test_unsupported_type_raises(self):
        f = SimpleUploadedFile("photo.png", b"\x89PNG", content_type="image/png")
        with self.assertRaises(UnsupportedFileType):
            imports.extract_emails_from_file(f)
